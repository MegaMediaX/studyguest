"""Extract course files into data/corpus/ and retrieve source chunks for a task.

    python -m studyquest.corpus                 # extract text (fast, no AI)
    python -m studyquest.corpus --ocr           # transcribe scanned pages (claude, cached, resumable)
    python -m studyquest.corpus --ocr --provider gemini --workers 3
    python -m studyquest.corpus --status        # what is extracted / scanned / OCR'd
"""
import argparse
import hashlib
import math
import re
import subprocess
import sys
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from . import ai, store

MIN_TEXT_CHARS = 40
CHUNK_CHARS = 2500
EXTS = {".pdf", ".pptx", ".docx", ".doc", ".md"}
OCR_PROMPT = (
    "Transcribe the course page in the image file {path} exactly. Keep headings, numbering "
    "(Example 3, Q2, 1.14...), formulas in plain-text math (e.g. f_x = 2xy, ∫∫_R, sqrt(x^2+y^2)). "
    "No commentary. If the page is blank, reply BLANK."
)


# ---------- extraction ----------

def _pdf_pages(path: Path) -> list[str]:
    import pymupdf
    with pymupdf.open(path) as doc:
        return [page.get_text() for page in doc]


def _pptx_pages(path: Path) -> list[str]:
    from pptx import Presentation
    slides = []
    for slide in Presentation(str(path)).slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                parts += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide:
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        slides.append("\n".join(p for p in parts if p.strip()))
    return slides


def _split(text: str) -> list[str]:
    paras, chunks, buf = text.split("\n"), [], ""
    for p in paras:
        if len(buf) + len(p) > CHUNK_CHARS and buf:
            chunks.append(buf)
            buf = ""
        buf += p + "\n"
    return chunks + ([buf] if buf.strip() else [])


def _docx_pages(path: Path) -> list[str]:
    import docx
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        lines += [" | ".join(c.text for c in row.cells) for row in table.rows]
    return _split("\n".join(lines))


def _doc_pages(path: Path) -> list[str]:
    out = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)], capture_output=True, text=True)
    return _split(out.stdout)


def _md_pages(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    sections = re.split(r"\n(?=#{1,3} )", text)
    return [c for s in sections for c in _split(s)]


EXTRACTORS = {".pdf": _pdf_pages, ".pptx": _pptx_pages, ".docx": _docx_pages, ".doc": _doc_pages, ".md": _md_pages}
UNIT = {".pdf": "page", ".pptx": "slide", ".docx": "part", ".doc": "part", ".md": "section"}


def course_files() -> list[tuple[str, Path, str]]:
    """(course, absolute path, path relative to the course folder)."""
    cfg = store.load_config()["courses"]
    out = []
    for code, c in cfg.items():
        folder = store.ROOT / c["folder"]
        if folder.exists():
            for p in sorted(folder.rglob("*")):
                if p.suffix.lower() in EXTS and p.is_file() and not p.name.startswith("~$"):
                    out.append((code, p, str(p.relative_to(folder))))
        notes = store.ROOT / c.get("notes", "")
        if c.get("notes") and notes.is_file():
            out.append((code, notes, notes.name))
    return out


def corpus_path(course: str, rel: str) -> Path:
    """Only configured course codes; the result is always inside data/corpus/."""
    if course not in store.load_config()["courses"]:
        raise KeyError(f"Unknown course {course!r}")
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", rel).lstrip(".")
    path = store.CORPUS / course / f"{slug}.json"
    if not path.resolve().is_relative_to(store.CORPUS.resolve()):
        raise KeyError("Bad source path")
    return path


def extract_all(verbose: bool = True) -> dict:
    store.ensure_dirs()
    summary = {"files": 0, "pages": 0, "scanned": 0, "errors": []}
    for course, path, rel in course_files():
        out = corpus_path(course, rel)
        mtime = path.stat().st_mtime
        old = store.read_json(out)
        if old and old.get("mtime") == mtime:
            _tally(summary, old)
            continue
        try:
            texts = EXTRACTORS[path.suffix.lower()](path)
        except Exception as e:  # noqa: BLE001 - one broken file must not stop the rest
            summary["errors"].append(f"{rel}: {e}")
            continue
        prev = {pg["n"]: pg for pg in (old or {}).get("pages", []) if pg["src"] == "ocr"}
        pages = []
        for i, t in enumerate(texts, 1):
            if len(t.strip()) >= MIN_TEXT_CHARS:
                pages.append({"n": i, "src": "text", "text": t.strip()})
            else:
                pages.append(prev.get(i) or {"n": i, "src": "none", "text": ""})
        doc = {"course": course, "file": rel, "path": str(path), "unit": UNIT[path.suffix.lower()],
               "mtime": mtime, "pages": pages}
        store.write_json(out, doc)
        _tally(summary, doc)
        if verbose:
            print(f"  {course}  {rel}: {len(pages)} {doc['unit']}s, {_count(doc, 'none')} without text")
    _INDEX.clear()
    return summary


def _count(doc: dict, src: str) -> int:
    return sum(1 for p in doc["pages"] if p["src"] == src)


def _tally(summary: dict, doc: dict) -> None:
    summary["files"] += 1
    summary["pages"] += len(doc["pages"])
    summary["scanned"] += _count(doc, "none")


def load_docs() -> list[dict]:
    return [store.read_json(p) for p in sorted(store.CORPUS.rglob("*.json"))]


def status_report() -> list[dict]:
    rows = []
    for d in load_docs():
        rows.append({"course": d["course"], "file": d["file"], "units": len(d["pages"]),
                     "text": _count(d, "text"), "ocr": _count(d, "ocr"), "no_text": _count(d, "none")})
    return rows


# ---------- OCR for scanned pages ----------

def _render(path: Path, n: int) -> Path | None:
    """PDF page -> PNG; PPTX slide -> its largest embedded picture. None if nothing to read."""
    png_dir = store.DATA / "ocr_png"
    png_dir.mkdir(parents=True, exist_ok=True)
    stem = hashlib.sha1(f"{path}|{n}".encode()).hexdigest()[:16]
    if path.suffix.lower() == ".pptx":
        from pptx import Presentation
        pics = [sh.image for sh in Presentation(str(path)).slides[n - 1].shapes if hasattr(sh, "image")]
        if not pics:
            return None
        img = max(pics, key=lambda im: len(im.blob))
        out = png_dir / f"{stem}.{img.ext}"
        if not out.exists():
            out.write_bytes(img.blob)
        return out
    import pymupdf
    out = png_dir / f"{stem}.png"
    if not out.exists():
        with pymupdf.open(path) as doc:
            doc[n - 1].get_pixmap(dpi=150).save(out)
    return out


def ocr_page(provider: str, path: Path, n: int) -> str:
    png = _render(path, n)
    if png is None:
        return ""
    text = ai.ask(provider, OCR_PROMPT.format(path=png), images=[png], want_json=False)
    return "" if text.strip().upper() == "BLANK" else text


def ocr_all(provider: str = "claude", workers: int = 3, limit: int | None = None) -> None:
    jobs, skip = [], set(store.load_config().get("ocr_skip", []))
    for d in load_docs():
        if not d["file"].lower().endswith((".pdf", ".pptx")) or d["file"] in skip:
            continue
        jobs += [(d, pg["n"]) for pg in d["pages"] if pg["src"] == "none"]
    jobs = jobs[:limit] if limit else jobs
    print(f"OCR: {len(jobs)} scanned pages via {provider}, {workers} at a time (cached, safe to stop and resume)")
    done = 0
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futs = {pool.submit(ocr_page, provider, Path(d["path"]), n): (d, n) for d, n in jobs}
        for fut in as_completed(futs):
            d, n = futs[fut]
            done += 1
            try:
                text = fut.result()
            except ai.AIUnavailable as e:
                print(f"  [{done}/{len(jobs)}] {d['file']} p{n}: FAILED ({e})")
                continue
            doc = store.read_json(corpus_path(d["course"], d["file"]))
            for pg in doc["pages"]:
                if pg["n"] == n:
                    pg.update(src="ocr" if text.strip() else "none", text=text.strip(), ocr_by=provider)
            store.write_json(corpus_path(d["course"], d["file"]), doc)
            print(f"  [{done}/{len(jobs)}] {d['file']} p{n}: {len(text)} chars")
    _INDEX.clear()


# ---------- retrieval ----------

TOKEN_RE = re.compile(r"\d+\.\d+|[a-zA-Zα-ωΑ-Ω∇]+|\d+")
STOP = set("the a an and or of to in on for with by at from is are be as it this that do read solve redo check "
           "sit write list pdf pptx notes ex q".split())
_INDEX: dict = {}


def tokens(text: str) -> list[str]:
    return [t.lower() for t in TOKEN_RE.findall(text) if t.lower() not in STOP and len(t) > 1 or "." in t]


def _index() -> dict:
    if not _INDEX:
        chunks = []
        for d in load_docs():
            for pg in d["pages"]:
                if pg["src"] != "none":
                    chunks.append({"course": d["course"], "file": d["file"], "unit": d["unit"], "n": pg["n"],
                                   "src": pg["src"], "text": pg["text"], "tf": Counter(tokens(pg["text"]))})
        df = Counter(t for c in chunks for t in c["tf"])
        _INDEX.update(chunks=chunks, df=df, n=max(len(chunks), 1))
    return _INDEX


def course_for_subject(subject: str) -> str | None:
    for code, c in store.load_config()["courses"].items():
        if subject.strip().lower() in {a.lower() for a in c["aliases"]} or subject.upper().startswith(code):
            return code
    return None


def file_hints(task: str) -> list[str]:
    """Filenames named in the task: quoted names, bare *.pdf/*.pptx, or configured aliases."""
    low = task.lower()
    hints = [q.strip() for q in re.findall(r"[“\"]([^”\"]+)[”\"]", task)]
    unquoted = re.sub(r"[“\"][^”\"]+[”\"]", " ", task)
    hints += re.findall(r"[\w\-.()]+\.(?:pdf|pptx|docx)", unquoted)
    for alias, fname in store.load_config().get("file_aliases", {}).items():
        if re.search(r"\b" + re.escape(alias), low):
            hints.append(fname)
    return list(dict.fromkeys(h.strip() for h in hints if h.strip()))  # dedupe, keep order


def unit_range(task: str) -> tuple[int, int] | None:
    m = re.search(r"(?:slides?|pp?\.?|pages?)\s*(\d+)\s*[–\-]\s*(\d+)", task, re.I)
    return (int(m.group(1)), int(m.group(2))) if m else None


def _match_file(chunk_file: str, hint: str) -> int:
    """2 = exact filename, 1 = loose match, 0 = no match."""
    base = Path(chunk_file).name.lower()
    h = hint.lower()
    if base == h or chunk_file.lower() == h:
        return 2
    stem = re.sub(r"\.(pdf|pptx|docx)$", "", h)
    # loose = the base name starts with the hint, or contains it after a separator ("Ch.2 X" in "Course Notes/Ch.2 X")
    return 1 if len(stem) > 3 and re.search(r"(?:^|[\s/_(\-])" + re.escape(stem), base) else 0


def retrieve(task: str, session: str, subject: str, k: int = 4) -> list[dict]:
    idx = _index()
    course = course_for_subject(subject)
    pool = [c for c in idx["chunks"] if course is None or c["course"] == course]

    hints = file_hints(task)
    if hints:
        scored = [(max(_match_file(c["file"], h) for h in hints), c) for c in pool]
        best = max((s for s, _ in scored), default=0)
        if not best:
            return []  # the task names a file we have no text for: don't quote some other file instead
        pool = [c for s, c in scored if s == best]

    rng = unit_range(task)
    if rng and hints:
        ranged = [c for c in pool if rng[0] <= c["n"] <= rng[1]]
        if ranged:
            step = max(len(ranged) // k, 1)  # spread questions across the whole assigned range
            return [_public(c) for c in ranged[::step][:k]]

    section = re.search(r"textbook\s+(\d+\.\d+)", task, re.I)
    query = tokens(task + " " + session) + ([section.group(1)] * 3 if section else [])
    ranked = sorted(pool, key=lambda c: _score(c, query, idx), reverse=True)
    return [_public(c) for c in ranked[:k] if _score(c, query, idx) > 0]


def retrieve_exam(task: str, session: str, subject: str, date: str | None = None, k: int = 3) -> list[dict]:
    """Past-exam pages for this task (config "exam_models": the papers of the next exam after the session's date),
    so battles are modelled on real exams. Best keyword matches first; the papers themselves if nothing matches."""
    course = course_for_subject(subject)
    groups = store.load_config().get("exam_models", {}).get(course or "", [])
    group = next((g for g in groups if (date or "") <= g["until"]), groups[-1] if groups else None)
    if not group:
        return []
    idx = _index()
    pool = [c for c in idx["chunks"] if c["course"] == course and any(_match_file(c["file"], f) for f in group["files"])]
    query = tokens(task + " " + session)
    ranked = sorted(pool, key=lambda c: (_score(c, query, idx), -group["files"].index(
        next(f for f in group["files"] if _match_file(c["file"], f)))), reverse=True)
    return [_public(c) for c in ranked[:k]]


def _score(c: dict, query: list[str], idx: dict) -> float:
    s = 0.0
    for t in query:
        if t in c["tf"]:
            s += (1 + math.log(c["tf"][t])) * math.log(1 + idx["n"] / idx["df"][t])
    return s


def _public(c: dict) -> dict:
    return {"file": c["file"], "unit": c["unit"], "n": c["n"], "src": c["src"], "course": c["course"],
            "text": c["text"][:CHUNK_CHARS]}


def missing_files_for(task: str, subject: str) -> list[str]:
    """Files the task names that aren't in the corpus at all (renamed, typo, not downloaded)."""
    course = course_for_subject(subject)
    docs = [d for d in load_docs() if not course or d["course"] == course]
    return [h for h in file_hints(task) if not any(_match_file(d["file"], h) for d in docs)]


def scanned_files_for(task: str, subject: str) -> list[str]:
    """Files the task names that exist but have no extractable text yet."""
    course = course_for_subject(subject)
    out = []
    for d in load_docs():
        if course and d["course"] != course:
            continue
        if any(_match_file(d["file"], h) for h in file_hints(task)) and _count(d, "text") + _count(d, "ocr") == 0:
            out.append(d["file"])
    return out


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ocr", action="store_true")
    ap.add_argument("--provider", default="claude", choices=["claude", "gemini"])
    ap.add_argument("--workers", type=int, default=3)
    ap.add_argument("--limit", type=int)
    ap.add_argument("--status", action="store_true")
    a = ap.parse_args(argv)
    if a.status:
        for r in status_report():
            flag = "  ⚠ no text" if r["no_text"] else ""
            print(f"{r['course']:8} {r['file'][:60]:60} {r['units']:5} text={r['text']} ocr={r['ocr']}{flag}")
        return
    s = extract_all()
    print(f"Extracted {s['files']} files, {s['pages']} pages/slides; {s['scanned']} have no text.")
    for e in s["errors"]:
        print("  ERROR", e, file=sys.stderr)
    if a.ocr:
        ocr_all(a.provider, a.workers, a.limit)


if __name__ == "__main__":
    main()
